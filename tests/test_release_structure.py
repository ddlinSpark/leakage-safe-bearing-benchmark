"""Tests for the formal paper-repository layout and result catalogs."""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]


def normalized_payload(path: Path) -> bytes:
    payload = path.read_bytes()
    if path.suffix.lower() in {".csv", ".json", ".md", ".txt"}:
        payload = payload.replace(b"\r\n", b"\n")
    return payload


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    digest.update(normalized_payload(path))
    return digest.hexdigest()


def test_readme_is_concise_and_contains_required_sections() -> None:
    text = (ROOT / "README.md").read_text(encoding="utf-8")
    assert len(text.splitlines()) <= 250
    for heading in (
        "## Key Findings",
        "## Generalization Gap",
        "## Experiment Status",
        "## Leakage-Safe Protocol",
        "## Quick Verification",
        "## Full Reproduction",
        "## Data Availability",
        "## Reproducibility",
        "## Scope and Limitations",
    ):
        assert heading in text
    assert "not all evaluated methods produced positive" in text


def test_formal_documentation_and_entry_points_exist() -> None:
    for relative in (
        "docs/DATASET.md",
        "docs/PROTOCOL.md",
        "docs/EXPERIMENTS.md",
        "docs/RESULTS.md",
        "docs/REPRODUCIBILITY.md",
        "docs/reports/FINAL_MANUSCRIPT_FORMAT_REPORT.md",
        "docs/reports/MANUSCRIPT_GITHUB_CONSISTENCY_REPORT.md",
        "docs/reports/REFERENCE_LINK_REPORT.md",
        "docs/reports/README_LINK_CHECK.md",
        "scripts/prepare_data.py",
        "scripts/run_main_benchmark.py",
        "scripts/generate_statistics.py",
        "scripts/generate_figures.py",
        "scripts/reproduce_all.py",
        "scripts/smoke_test.py",
        "scripts/audit_internal_links.py",
        "src/spindle_cnn/visualization/style.py",
    ):
        assert (ROOT / relative).is_file(), relative
    report = (ROOT / "docs/reports/REFERENCE_LINK_REPORT.md").read_text(
        encoding="utf-8"
    )
    assert "**PASS**" in report
    assert "pdf_bookmarks: PASS" in report


def test_root_cleanup_and_pdf_free_release() -> None:
    root_markdown = {
        path.name for path in ROOT.glob("*.md")
    }
    assert root_markdown <= {
        "README.md",
        "FINAL_RELEASE_REPORT.md",
    }
    assert not list((ROOT / "figures" / "publication").glob("*.pdf"))
    assert not list((ROOT / "manuscript").glob("*.pdf"))
    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    pyproject = (ROOT / "pyproject.toml").read_text(encoding="utf-8")
    assert "manuscript/manuscript.pdf" not in pyproject
    image_path = ROOT / "figures/publication/figure_09_validation_test_gap.png"
    assert image_path.is_file()
    assert "figure_09_validation_test_gap.png" in readme


def test_final_manuscript_format_and_availability() -> None:
    document = Document(
        ROOT / "docs/paper/Leakage_Safe_Bearing_Benchmark_SCI_final.docx"
    )
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    text = "\n".join(paragraphs + cells)
    assert "+/-" not in text
    assert re.search(r"(?<=[0-9]) x (?=[0-9])|folds x [0-9]", text) is None
    assert re.search(r"(?i)(?<![A-Za-z])tau(?![A-Za-z])", text) is None
    assert "Appendix A" not in text
    assert "Supplementary Material" in text
    assert "Canonical completed evaluations" in text
    assert "Held-out test" in text
    assert "https://github.com/ddlinSpark/leakage-safe-bearing-benchmark" in text
    assert document.core_properties.author == "Huadong Lin"


def test_canonical_and_split_manifests_match_files() -> None:
    canonical = json.loads(
        (ROOT / "results" / "canonical" / "manifest.json").read_text(
            encoding="utf-8"
        )
    )
    for relative, record in canonical["files"].items():
        path = ROOT / relative
        assert len(normalized_payload(path)) == record["bytes"], relative
        assert sha256(path) == record["sha256"], relative

    splits = json.loads(
        (ROOT / "splits" / "manifest.json").read_text(encoding="utf-8")
    )
    assert splits["authoritative_directory"] == "results/splits"
    for relative, record in splits["files"].items():
        path = ROOT / relative
        assert len(normalized_payload(path)) == record["bytes"], relative
        assert sha256(path) == record["sha256"], relative
