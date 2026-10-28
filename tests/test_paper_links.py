"""Tests for internal citation and cross-reference links."""

from __future__ import annotations

from docx import Document
from docx.oxml.ns import qn

from spindle_cnn.paper.docx_style import (
    add_caption,
    add_heading_bookmarks,
    add_inline_text,
    add_reference,
    configure_document,
)


def anchors(paragraph) -> list[str]:
    return [
        value
        for value in paragraph._p.xpath(".//w:hyperlink/@w:anchor")
    ]


def bookmark_names(document: Document) -> set[str]:
    return {
        value
        for value in document.element.body.xpath(
            ".//w:bookmarkStart/@w:name"
        )
    }


def test_inline_text_links_reference_numbers_and_cross_references() -> None:
    document = Document()
    configure_document(document)
    paragraph = document.add_paragraph()
    add_inline_text(
        paragraph,
        "See [1], [2,3], [11-13], Figure 4, Table 13, Section 6, and Eq. (2).",
    )
    assert anchors(paragraph) == [
        "ref_1",
        "ref_2",
        "ref_3",
        "ref_11",
        "ref_13",
        "fig_4",
        "tbl_13",
        "sec_6",
        "eq_2",
    ]
    assert paragraph.text == (
        "See [1], [2,3], [11-13], Figure 4, Table 13, Section 6, and Eq. (2)."
    )


def test_inline_text_normalizes_manuscript_notation_only() -> None:
    document = Document()
    configure_document(document)
    paragraph = document.add_paragraph()
    add_inline_text(
        paragraph,
        "4 folds x 5 seeds; 96 x 96; Conv5x5; 0.289 +/- 0.104; "
        "tau = 0.60; -0.020; macro-F1.",
    )
    assert paragraph.text == (
        "4 folds × 5 seeds; 96 × 96; Conv5x5; 0.289 ± 0.104; "
        "τ = 0.60; −0.020; Macro-F1."
    )


def test_caption_heading_and_reference_bookmarks_are_created() -> None:
    document = Document()
    configure_document(document)
    add_caption(document, "Figure 4. Main comparison.", above=False)
    add_caption(document, "Table 13. Generalization gap.", above=True)
    document.add_heading("3.2 Data and protocol", level=2)
    document.add_heading("Supplementary Material", level=1)
    document.add_heading("Additional Diagnostic Figures", level=2)
    document.add_heading("References", level=1)
    add_reference(document, 1, "C. Lessmeier et al.")
    add_heading_bookmarks(document)
    names = bookmark_names(document)
    assert {
        "fig_4",
        "tbl_13",
        "sec_3_2",
        "supplementary",
        "refs",
        "ref_1",
    }.issubset(names)
