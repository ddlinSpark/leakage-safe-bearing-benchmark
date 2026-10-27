"""Independent pre-submission audit of results and manuscript tables."""

from __future__ import annotations

import argparse
import json
import math
from itertools import product
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from scipy import stats

from docx import Document


METHOD_LABELS = {
    "raw_cnn": "Raw 1D-CNN",
    "gru_only": "GRU only",
    "cwt_cnn": "CWT 2D-CNN",
    "cwt_gru_no_se": "CWT-GRU without SE",
    "cwt_gru": "CWT-GRU",
    "cwt_gru_temp": "CWT-GRU + temperature scaling",
    "cwt_gru_selftrain": "CWT-GRU + temperature scaling + self-training",
}

METHOD_ORDER = list(METHOD_LABELS)
SEEDS = [42, 123, 2024, 2025, 3407]
FOLDS = [0, 1, 2, 3]
TOL = 1e-9


def f3(value: float) -> str:
    return f"{float(value):.3f}"


def p_text(value: float) -> str:
    return "< 0.001" if float(value) < 0.001 else f"= {float(value):.3f}"


def percent(value: float) -> str:
    return f"{100 * float(value):.1f}%"


def mean_sd(mean: float, std: float) -> str:
    return f"{float(mean):.3f} +/- {float(std):.3f}"


def ci95(values: pd.Series) -> tuple[float, float]:
    clean = values.replace([np.inf, -np.inf], np.nan).dropna().to_numpy(float)
    if len(clean) < 2:
        return float("nan"), float("nan")
    half_width = stats.t.ppf(0.975, len(clean) - 1) * clean.std(ddof=1) / math.sqrt(
        len(clean)
    )
    return float(clean.mean() - half_width), float(clean.mean() + half_width)


def independent_hierarchical_values(
    raw: pd.DataFrame,
    external: pd.DataFrame,
    row: pd.Series,
) -> tuple[float, float, float, float]:
    reference = str(row["reference"])
    method = str(row["method"])
    experiment_type = str(row["experiment_type"])
    if experiment_type == "main":
        selected = raw.loc[
            raw["experiment_type"].eq("main")
            & raw["method"].isin([reference, method])
        ]
    elif experiment_type == "label_efficiency":
        selected = raw.loc[
            raw["experiment_type"].eq("label_efficiency")
            & np.isclose(
                raw["label_ratio"].to_numpy(dtype=float),
                float(row["label_ratio"]),
            )
            & raw["method"].isin([reference, method])
        ]
    elif experiment_type == "label_full":
        selected = raw.loc[
            raw["experiment_type"].eq("main")
            & np.isclose(raw["label_ratio"].to_numpy(dtype=float), 1.0)
            & raw["method"].isin([reference, method])
        ]
    elif experiment_type == "external":
        selected = pd.concat(
            [
                raw.loc[
                    raw["experiment_type"].eq("main")
                    & raw["method"].eq(reference)
                ],
                external.loc[external["method"].eq(method)],
            ],
            ignore_index=True,
        )
    else:
        raise ValueError(f"Unknown hierarchical experiment type: {experiment_type}")

    pivot = selected.pivot_table(
        index=["fold", "seed"],
        columns="method",
        values="macro_f1",
    )
    paired = pivot[[reference, method]].dropna()
    expected = {
        (fold, seed) for fold in FOLDS for seed in SEEDS
    }
    if set(paired.index) != expected:
        raise ValueError(
            f"Incomplete hierarchy input for {row['comparison']}"
        )
    differences = paired[method] - paired[reference]
    values_by_fold = {
        fold: differences.loc[fold].to_numpy(dtype=float) for fold in FOLDS
    }
    rng = np.random.default_rng(int(row["bootstrap_seed"]))
    samples = int(row["bootstrap_samples"])
    bootstrap_means = np.empty(samples, dtype=float)
    for index in range(samples):
        selected_folds = rng.choice(FOLDS, size=len(FOLDS), replace=True)
        selected_values = []
        for fold in selected_folds:
            fold_values = values_by_fold[int(fold)]
            seed_indices = rng.integers(
                0,
                len(fold_values),
                size=len(fold_values),
            )
            selected_values.append(fold_values[seed_indices])
        bootstrap_means[index] = np.concatenate(selected_values).mean()
    low, high = np.quantile(bootstrap_means, [0.025, 0.975])
    fold_means = differences.groupby(level="fold").mean().to_numpy(dtype=float)
    observed = float(fold_means.mean())
    if np.allclose(fold_means, 0.0, rtol=0.0, atol=1e-15):
        fold_p = 1.0
    else:
        permuted = [
            float(np.mean(fold_means * np.asarray(signs)))
            for signs in product((-1.0, 1.0), repeat=len(fold_means))
        ]
        fold_p = float(
            np.sum(np.abs(permuted) >= abs(observed) - 1e-15)
            / len(permuted)
        )
    return float(differences.mean()), float(low), float(high), fold_p


def check_hierarchical_statistics(
    raw: pd.DataFrame,
    external: pd.DataFrame,
    statistics: pd.DataFrame,
    mismatches: list[dict[str, Any]],
) -> None:
    if not statistics["comparison"].is_unique:
        add_mismatch(
            mismatches,
            location="hierarchical statistics",
            paper_value="unique comparison identifiers",
            computed_value="duplicate comparison identifiers",
            source="results/hierarchical_statistics.csv",
            reason="Comparison keys must be unique.",
        )
    for _, row in statistics.iterrows():
        mean, low, high, fold_p = independent_hierarchical_values(
            raw,
            external,
            row,
        )
        location = f"hierarchical comparison: {row['comparison']}"
        assert_close(
            float(row["mean_delta"]),
            mean,
            location=f"{location}, mean delta",
            source="results/raw_runs.csv and results/sota/raw_runs.csv",
            mismatches=mismatches,
        )
        assert_close(
            float(row["hierarchical_ci_low"]),
            low,
            location=f"{location}, lower interval",
            source="results/raw_runs.csv and results/sota/raw_runs.csv",
            mismatches=mismatches,
        )
        assert_close(
            float(row["hierarchical_ci_high"]),
            high,
            location=f"{location}, upper interval",
            source="results/raw_runs.csv and results/sota/raw_runs.csv",
            mismatches=mismatches,
        )
        assert_close(
            float(row["fold_aware_p"]),
            fold_p,
            location=f"{location}, fold sign-flip p-value",
            source="results/raw_runs.csv and results/sota/raw_runs.csv",
            mismatches=mismatches,
        )
        if (
            int(row["n_folds"]) != 4
            or int(row["n_seeds"]) != 5
            or int(row["n_pairs"]) != 20
        ):
            add_mismatch(
                mismatches,
                location=f"{location}, design dimensions",
                paper_value="4 folds, 5 seeds, 20 pairs",
                computed_value=(
                    f"{int(row['n_folds'])} folds, "
                    f"{int(row['n_seeds'])} seeds, "
                    f"{int(row['n_pairs'])} pairs"
                ),
                source="results/hierarchical_statistics.csv",
                reason="Fold-seed completeness differs from the protocol.",
            )


def add_mismatch(
    mismatches: list[dict[str, Any]],
    *,
    location: str,
    paper_value: str,
    computed_value: str,
    source: str,
    reason: str,
) -> None:
    mismatches.append(
        {
            "paper_location": location,
            "paper_value": paper_value,
            "computed_value": computed_value,
            "source": source,
            "reason": reason,
        }
    )


def assert_close(
    actual: float,
    expected: float,
    *,
    location: str,
    source: str,
    mismatches: list[dict[str, Any]],
) -> None:
    if not np.isclose(actual, expected, rtol=0, atol=TOL, equal_nan=True):
        add_mismatch(
            mismatches,
            location=location,
            paper_value=f"{expected:.12g}",
            computed_value=f"{actual:.12g}",
            source=source,
            reason="Numerical value differs from independent recomputation.",
        )


def compare_table(
    actual: list[list[str]],
    expected: list[list[str]],
    *,
    location: str,
    source: str,
    mismatches: list[dict[str, Any]],
) -> None:
    if len(actual) != len(expected):
        add_mismatch(
            mismatches,
            location=f"{location}, row count",
            paper_value=str(len(expected)),
            computed_value=str(len(actual)),
            source=source,
            reason="Table row count differs.",
        )
    for row_index, (actual_row, expected_row) in enumerate(
        zip(actual, expected, strict=False),
        start=1,
    ):
        if len(actual_row) != len(expected_row):
            add_mismatch(
                mismatches,
                location=f"{location}, row {row_index}, column count",
                paper_value=str(len(expected_row)),
                computed_value=str(len(actual_row)),
                source=source,
                reason="Table column count differs.",
            )
            continue
        for column_index, (actual_value, expected_value) in enumerate(
            zip(actual_row, expected_row, strict=True),
            start=1,
        ):
            if str(actual_value).strip() != str(expected_value).strip():
                add_mismatch(
                    mismatches,
                    location=f"{location}, row {row_index}, column {column_index}",
                    paper_value=str(expected_value),
                    computed_value=str(actual_value),
                    source=source,
                    reason="Paper table value differs from independently computed value.",
                )


def table_rows(table) -> list[list[str]]:
    return [[cell.text.strip() for cell in row.cells] for row in table.rows]


def independently_summarize(frame: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for keys, group in frame.groupby(
        ["experiment_type", "method", "label_ratio", "threshold"],
        dropna=False,
    ):
        experiment_type, method, ratio, threshold = keys
        macro_low, macro_high = ci95(group["macro_f1"])
        accuracy_low, accuracy_high = ci95(group["accuracy"])
        ece_low, ece_high = ci95(group["ece"])
        rows.append(
            {
                "experiment_type": experiment_type,
                "method": method,
                "label_ratio": ratio,
                "threshold": threshold,
                "Runs": len(group),
                "Macro-F1 Mean": group["macro_f1"].mean(),
                "Macro-F1 Std": group["macro_f1"].std(ddof=1),
                "Macro-F1 CI95 Low": macro_low,
                "Macro-F1 CI95 High": macro_high,
                "Accuracy Mean": group["accuracy"].mean(),
                "Accuracy Std": group["accuracy"].std(ddof=1),
                "Accuracy CI95 Low": accuracy_low,
                "Accuracy CI95 High": accuracy_high,
                "ECE Mean": group["ece"].mean(),
                "ECE Std": group["ece"].std(ddof=1),
                "ECE CI95 Low": ece_low,
                "ECE CI95 High": ece_high,
            }
        )
    return pd.DataFrame(rows)


def independent_pairs(frame: pd.DataFrame) -> pd.DataFrame:
    work = frame.copy()
    work["_threshold"] = work["threshold"].where(
        work["experiment_type"].eq("threshold"),
        np.nan,
    )
    rows: list[dict[str, Any]] = []
    for (experiment_type, ratio, threshold), group in work.groupby(
        ["experiment_type", "label_ratio", "_threshold"],
        dropna=False,
    ):
        pivot = group.pivot_table(
            index=["fold", "seed"],
            columns="method",
            values="macro_f1",
            aggfunc="mean",
        )
        present = [method for method in METHOD_ORDER if method in pivot.columns]
        present.extend(sorted(set(pivot.columns) - set(present)))
        for position, reference in enumerate(present):
            for method in present[position + 1 :]:
                paired = pivot[[reference, method]].dropna()
                differences = paired[method] - paired[reference]
                if np.allclose(differences, 0):
                    p_value = 1.0
                    test = "all_differences_zero"
                else:
                    p_value = float(
                        stats.wilcoxon(
                            differences,
                            alternative="two-sided",
                            zero_method="wilcox",
                        ).pvalue
                    )
                    test = "wilcoxon_signed_rank"
                low, high = ci95(differences)
                rows.append(
                    {
                        "experiment_type": experiment_type,
                        "label_ratio": ratio,
                        "threshold": threshold,
                        "reference": reference,
                        "method": method,
                        "paired_runs": len(paired),
                        "reference_mean_macro_f1": paired[reference].mean(),
                        "method_mean_macro_f1": paired[method].mean(),
                        "mean_paired_delta": differences.mean(),
                        "std_paired_delta": differences.std(ddof=1),
                        "paired_delta_ci95_low": low,
                        "paired_delta_ci95_high": high,
                        "test": test,
                        "p_value": p_value,
                    }
                )
    return pd.DataFrame(rows)


def frame_with_label_100(raw: pd.DataFrame) -> pd.DataFrame:
    main = raw.loc[raw["status"].eq("success")].copy()
    label_100 = main.loc[
        main["experiment_type"].eq("main")
        & main["label_ratio"].eq(1.0)
        & main["method"].isin(["cwt_gru", "cwt_gru_selftrain"])
    ].copy()
    label_100["experiment_type"] = "label_efficiency"
    label_100["run_id"] = label_100["run_id"].astype(str) + "__label100"
    return pd.concat([main, label_100], ignore_index=True).drop_duplicates("run_id")


def check_summary_sources(
    raw: pd.DataFrame,
    summary: pd.DataFrame,
    summary_ci: pd.DataFrame,
    pairs: pd.DataFrame,
    label_summary: pd.DataFrame,
    threshold_summary: pd.DataFrame,
    ablation_summary: pd.DataFrame,
    model_complexity: pd.DataFrame,
    mismatches: list[dict[str, Any]],
) -> pd.DataFrame:
    frame = frame_with_label_100(raw)
    computed = independently_summarize(frame)
    computed_main = computed.loc[computed["experiment_type"].eq("main")]
    if len(summary) != len(computed_main):
        add_mismatch(
            mismatches,
            location="summary.csv row count",
            paper_value=str(len(computed_main)),
            computed_value=str(len(summary)),
            source="results/summary.csv",
            reason="Main-model summary is incomplete or duplicated.",
        )
    for _, row in computed_main.iterrows():
        match = summary.loc[summary["method"].eq(row["method"])]
        location = f"summary.csv {row['method']}"
        if len(match) != 1:
            add_mismatch(
                mismatches,
                location=location,
                paper_value="one matching row",
                computed_value=f"{len(match)} matching rows",
                source="results/summary.csv",
                reason="Main-model summary row is missing or duplicated.",
            )
            continue
        actual = match.iloc[0]
        for column in (
            "Runs",
            "Macro-F1 Mean",
            "Macro-F1 Std",
            "Macro-F1 CI95 Low",
            "Macro-F1 CI95 High",
            "Accuracy Mean",
            "Accuracy Std",
            "Accuracy CI95 Low",
            "Accuracy CI95 High",
            "ECE Mean",
            "ECE Std",
            "ECE CI95 Low",
            "ECE CI95 High",
        ):
            assert_close(
                float(actual[column]),
                float(row[column]),
                location=f"{location}, {column}",
                source="results/summary.csv",
                mismatches=mismatches,
            )
    for _, row in computed.iterrows():
        match = summary_ci.loc[
            summary_ci["experiment_type"].eq(row["experiment_type"])
            & summary_ci["method"].eq(row["method"])
            & summary_ci["label_ratio"].round(6).eq(round(row["label_ratio"], 6))
            & (
                summary_ci["threshold"].isna()
                if pd.isna(row["threshold"])
                else summary_ci["threshold"].round(6).eq(round(row["threshold"], 6))
            )
        ]
        location = (
            f"summary_with_ci.csv {row['experiment_type']}/{row['method']}/"
            f"{row['label_ratio']}/{row['threshold']}"
        )
        if len(match) != 1:
            add_mismatch(
                mismatches,
                location=location,
                paper_value="one matching row",
                computed_value=f"{len(match)} matching rows",
                source="results/summary_with_ci.csv",
                reason="Summary row is missing or duplicated.",
            )
            continue
        actual = match.iloc[0]
        for column in (
            "Runs",
            "Macro-F1 Mean",
            "Macro-F1 Std",
            "Macro-F1 CI95 Low",
            "Macro-F1 CI95 High",
            "Accuracy Mean",
            "Accuracy Std",
            "Accuracy CI95 Low",
            "Accuracy CI95 High",
            "ECE Mean",
            "ECE Std",
            "ECE CI95 Low",
            "ECE CI95 High",
        ):
            assert_close(
                float(actual[column]),
                float(row[column]),
                location=f"{location}, {column}",
                source="results/summary_with_ci.csv",
                mismatches=mismatches,
            )

    computed_pairs = independent_pairs(frame)
    if len(computed_pairs) != len(pairs):
        add_mismatch(
            mismatches,
            location="statistical_tests.csv row count",
            paper_value=str(len(computed_pairs)),
            computed_value=str(len(pairs)),
            source="results/statistical_tests.csv",
            reason="Paired-test matrix is incomplete or duplicated.",
        )
    for _, row in computed_pairs.iterrows():
        match = pairs.loc[
            pairs["experiment_type"].eq(row["experiment_type"])
            & pairs["reference"].eq(row["reference"])
            & pairs["method"].eq(row["method"])
            & pairs["label_ratio"].round(6).eq(round(row["label_ratio"], 6))
            & (
                pairs["threshold"].isna()
                if pd.isna(row["threshold"])
                else pairs["threshold"].round(6).eq(round(row["threshold"], 6))
            )
        ]
        location = (
            f"statistical_tests.csv {row['experiment_type']}/"
            f"{row['reference']} vs {row['method']}/{row['label_ratio']}/"
            f"{row['threshold']}"
        )
        if len(match) != 1:
            add_mismatch(
                mismatches,
                location=location,
                paper_value="one matching row",
                computed_value=f"{len(match)} matching rows",
                source="results/statistical_tests.csv",
                reason="Paired-test row is missing or duplicated.",
            )
            continue
        actual = match.iloc[0]
        for column in (
            "paired_runs",
            "reference_mean_macro_f1",
            "method_mean_macro_f1",
            "mean_paired_delta",
            "std_paired_delta",
            "paired_delta_ci95_low",
            "paired_delta_ci95_high",
            "p_value",
        ):
            assert_close(
                float(actual[column]),
                float(row[column]),
                location=f"{location}, {column}",
                source="results/statistical_tests.csv",
                mismatches=mismatches,
            )
        if str(actual["test"]) != str(row["test"]):
            add_mismatch(
                mismatches,
                location=f"{location}, test",
                paper_value=str(row["test"]),
                computed_value=str(actual["test"]),
                source="results/statistical_tests.csv",
                reason="Recorded statistical test differs.",
            )

    computed_label = computed.loc[
        computed["experiment_type"].eq("label_efficiency")
    ]
    if len(computed_label) != len(label_summary):
        add_mismatch(
            mismatches,
            location="label_efficiency_summary.csv row count",
            paper_value=str(len(computed_label)),
            computed_value=str(len(label_summary)),
            source="results/label_efficiency_summary.csv",
            reason="Label-efficiency summary is incomplete or duplicated.",
        )
    for _, row in computed_label.iterrows():
        match = label_summary.loc[
            label_summary["method"].eq(row["method"])
            & label_summary["label_ratio"].round(6).eq(round(row["label_ratio"], 6))
        ]
        if len(match) != 1:
            add_mismatch(
                mismatches,
                location=(
                    "label_efficiency_summary.csv "
                    f"{row['method']}/{row['label_ratio']}"
                ),
                paper_value="one matching row",
                computed_value=f"{len(match)} matching rows",
                source="results/label_efficiency_summary.csv",
                reason="Label-efficiency row is missing or duplicated.",
            )
            continue
        actual = match.iloc[0]
        for column in (
            "Runs",
            "Macro-F1 Mean",
            "Macro-F1 Std",
            "Macro-F1 CI95 Low",
            "Macro-F1 CI95 High",
            "Accuracy Mean",
            "Accuracy Std",
            "Accuracy CI95 Low",
            "Accuracy CI95 High",
            "ECE Mean",
            "ECE Std",
            "ECE CI95 Low",
            "ECE CI95 High",
        ):
            assert_close(
                float(actual[column]),
                float(row[column]),
                location=(
                    "label_efficiency_summary.csv "
                    f"{row['method']}/{row['label_ratio']}/{column}"
                ),
                source="results/label_efficiency_summary.csv",
                mismatches=mismatches,
            )

    threshold = raw.loc[raw["experiment_type"].eq("threshold")]
    computed_threshold = (
        threshold.groupby(["method", "threshold"], as_index=False)
        .agg(
            macro_mean=("macro_f1", "mean"),
            macro_std=("macro_f1", "std"),
            accuracy_mean=("accuracy", "mean"),
            accuracy_std=("accuracy", "std"),
            ece_mean=("ece", "mean"),
            ece_std=("ece", "std"),
            pseudo_count_mean=("pseudo_count", "mean"),
            pseudo_count_std=("pseudo_count", "std"),
            pseudo_coverage_mean=("pseudo_coverage", "mean"),
            pseudo_coverage_std=("pseudo_coverage", "std"),
            pseudo_accuracy_mean=("pseudo_accuracy", "mean"),
            pseudo_accuracy_std=("pseudo_accuracy", "std"),
            runs=("seed", "count"),
        )
    )
    if len(computed_threshold) != len(threshold_summary):
        add_mismatch(
            mismatches,
            location="threshold_summary.csv row count",
            paper_value=str(len(computed_threshold)),
            computed_value=str(len(threshold_summary)),
            source="results/threshold_summary.csv",
            reason="Threshold summary is incomplete or duplicated.",
        )
    threshold_columns = {
        "Macro-F1 Mean": "macro_mean",
        "Macro-F1 Std": "macro_std",
        "Accuracy Mean": "accuracy_mean",
        "Accuracy Std": "accuracy_std",
        "ECE Mean": "ece_mean",
        "ECE Std": "ece_std",
        "pseudo_count_mean": "pseudo_count_mean",
        "pseudo_count_std": "pseudo_count_std",
        "pseudo_coverage_mean": "pseudo_coverage_mean",
        "pseudo_coverage_std": "pseudo_coverage_std",
        "pseudo_accuracy_mean": "pseudo_accuracy_mean",
        "pseudo_accuracy_std": "pseudo_accuracy_std",
        "Runs": "runs",
    }
    for _, row in computed_threshold.iterrows():
        match = threshold_summary.loc[
            threshold_summary["threshold"].round(6).eq(round(row["threshold"], 6))
        ]
        if len(match) != 1:
            add_mismatch(
                mismatches,
                location=f"threshold_summary.csv threshold={row['threshold']}",
                paper_value="one matching row",
                computed_value=f"{len(match)} matching rows",
                source="results/threshold_summary.csv",
                reason="Threshold summary row is missing or duplicated.",
            )
            continue
        actual = match.iloc[0]
        for actual_column, computed_column in threshold_columns.items():
            assert_close(
                float(actual[actual_column]),
                float(row[computed_column]),
                location=(
                    f"threshold_summary.csv threshold={row['threshold']}, "
                    f"{actual_column}"
                ),
                source="results/threshold_summary.csv",
                mismatches=mismatches,
            )

    components = {
        "cwt_cnn": (1, 0, 1, 0, 0),
        "gru_only": (0, 1, 0, 0, 0),
        "cwt_gru_no_se": (1, 1, 0, 0, 0),
        "cwt_gru": (1, 1, 1, 0, 0),
        "cwt_gru_temp": (1, 1, 1, 1, 0),
        "cwt_gru_selftrain": (1, 1, 1, 1, 1),
    }
    main = raw.loc[raw["experiment_type"].eq("main")]
    if len(components) != len(ablation_summary):
        add_mismatch(
            mismatches,
            location="ablation_summary.csv row count",
            paper_value=str(len(components)),
            computed_value=str(len(ablation_summary)),
            source="results/ablation_summary.csv",
            reason="Ablation summary is incomplete or duplicated.",
        )
    for method, values in components.items():
        group = main.loc[main["method"].eq(method)]
        match = ablation_summary.loc[ablation_summary["Model"].eq(method)]
        if len(match) != 1:
            add_mismatch(
                mismatches,
                location=f"ablation_summary.csv {method}",
                paper_value="one matching row",
                computed_value=f"{len(match)} matching rows",
                source="results/ablation_summary.csv",
                reason="Ablation row is missing or duplicated.",
            )
            continue
        actual = match.iloc[0]
        for column, expected_value in zip(
            ("CWT", "GRU", "SE", "Calibration", "Self-training"),
            values,
            strict=True,
        ):
            assert_close(
                float(actual[column]),
                float(expected_value),
                location=f"ablation_summary.csv {method}/{column}",
                source="results/ablation_summary.csv",
                mismatches=mismatches,
            )
        for column, value in (
            ("Runs", len(group)),
            ("Macro-F1 Mean", group["macro_f1"].mean()),
            ("Macro-F1 Std", group["macro_f1"].std(ddof=1)),
            ("Accuracy Mean", group["accuracy"].mean()),
            ("Accuracy Std", group["accuracy"].std(ddof=1)),
        ):
            assert_close(
                float(actual[column]),
                float(value),
                location=f"ablation_summary.csv {method}/{column}",
                source="results/ablation_summary.csv",
                mismatches=mismatches,
            )

    computed_complexity = (
        raw.groupby(["method", "model_key"], as_index=False)
        .agg(
            parameters=("parameters", "median"),
            gflops_sample=("gflops_sample", "median"),
            train_time_s_median=("train_time_s", "median"),
            inference_ms_per_batch_median=("inference_ms_per_batch", "median"),
            runs=("seed", "count"),
        )
    )
    if len(computed_complexity) != len(model_complexity):
        add_mismatch(
            mismatches,
            location="model_complexity.csv row count",
            paper_value=str(len(computed_complexity)),
            computed_value=str(len(model_complexity)),
            source="results/model_complexity.csv",
            reason="Model-complexity summary is incomplete or duplicated.",
        )
    for _, row in computed_complexity.iterrows():
        match = model_complexity.loc[model_complexity["method"].eq(row["method"])]
        if len(match) != 1:
            add_mismatch(
                mismatches,
                location=f"model_complexity.csv {row['method']}",
                paper_value="one matching row",
                computed_value=f"{len(match)} matching rows",
                source="results/model_complexity.csv",
                reason="Complexity row is missing or duplicated.",
            )
            continue
        actual = match.iloc[0]
        for column in (
            "parameters",
            "gflops_sample",
            "train_time_s_median",
            "inference_ms_per_batch_median",
            "runs",
        ):
            assert_close(
                float(actual[column]),
                float(row[column]),
                location=f"model_complexity.csv {row['method']}/{column}",
                source="results/model_complexity.csv",
                mismatches=mismatches,
            )
    return computed


def check_run_completeness(
    raw: pd.DataFrame,
    results_dir: Path,
    mismatches: list[dict[str, Any]],
) -> dict[str, Any]:
    expected = {
        "main": 140,
        "label_efficiency": 80,
        "threshold": 100,
    }
    counts = raw.groupby("experiment_type").size().to_dict()
    for experiment_type, count in expected.items():
        if int(counts.get(experiment_type, 0)) != count:
            add_mismatch(
                mismatches,
                location=f"raw_runs.csv {experiment_type} count",
                paper_value=str(count),
                computed_value=str(counts.get(experiment_type, 0)),
                source="results/raw_runs.csv",
                reason="Experiment run count differs from the declared matrix.",
            )

    key_columns = [
        "experiment_type",
        "method",
        "fold",
        "seed",
        "label_ratio",
        "threshold",
    ]
    duplicates = raw.loc[raw.duplicated(key_columns, keep=False), key_columns]
    if not duplicates.empty:
        add_mismatch(
            mismatches,
            location="raw_runs.csv experiment keys",
            paper_value="unique experiment keys",
            computed_value=f"{len(duplicates)} duplicate rows",
            source="results/raw_runs.csv",
            reason="Duplicate fold/seed/model/configuration keys exist.",
        )

    if not raw["status"].eq("success").all():
        add_mismatch(
            mismatches,
            location="raw_runs.csv status",
            paper_value="all success",
            computed_value=raw["status"].value_counts().to_dict(),
            source="results/raw_runs.csv",
            reason="Failed runs remain in the canonical table.",
        )

    required_non_null = [
        "fold",
        "seed",
        "method",
        "label_ratio",
        "macro_f1",
        "accuracy",
        "raw_macro_f1",
        "raw_accuracy",
        "ece",
        "nll",
        "brier",
        "parameters",
        "gflops_sample",
        "train_time",
        "inference_ms_per_batch",
        "labeled_windows",
        "unlabeled_candidate_windows",
        "train_windows",
        "validation_windows",
        "test_windows",
        "dataset_hash",
        "split_hash",
        "config_hash",
        "git_commit",
    ]
    null_counts = {
        column: int(raw[column].isna().sum())
        for column in required_non_null
        if raw[column].isna().any()
    }
    if null_counts:
        add_mismatch(
            mismatches,
            location="raw_runs.csv required fields",
            paper_value="no missing values",
            computed_value=null_counts,
            source="results/raw_runs.csv",
            reason="One or more required result fields are missing.",
        )

    main = raw.loc[raw["experiment_type"].eq("main")]
    expected_main = pd.MultiIndex.from_product(
        [METHOD_ORDER, FOLDS, SEEDS],
        names=["method", "fold", "seed"],
    )
    observed_main = pd.MultiIndex.from_frame(main[["method", "fold", "seed"]])
    if len(observed_main) != 140 or set(observed_main) != set(expected_main):
        add_mismatch(
            mismatches,
            location="main experiment matrix",
            paper_value="7 methods x 4 folds x 5 seeds",
            computed_value=f"{len(observed_main)} unique keys",
            source="results/raw_runs.csv",
            reason="Main fold/seed/method combinations are incomplete.",
        )

    label = raw.loc[raw["experiment_type"].eq("label_efficiency")]
    expected_label = pd.MultiIndex.from_product(
        [["cwt_gru", "cwt_gru_selftrain"], [0.1, 0.5], FOLDS, SEEDS],
        names=["method", "label_ratio", "fold", "seed"],
    )
    observed_label = pd.MultiIndex.from_frame(
        label[["method", "label_ratio", "fold", "seed"]]
    )
    if len(observed_label) != 80 or set(observed_label) != set(expected_label):
        add_mismatch(
            mismatches,
            location="label-efficiency matrix",
            paper_value="2 methods x 2 saved budgets x 4 folds x 5 seeds",
            computed_value=f"{len(observed_label)} unique keys",
            source="results/raw_runs.csv",
            reason="Label-efficiency combinations are incomplete.",
        )

    threshold = raw.loc[raw["experiment_type"].eq("threshold")]
    thresholds = [0.3, 0.4, 0.5, 0.6, 0.7]
    expected_threshold = pd.MultiIndex.from_product(
        [thresholds, FOLDS, SEEDS],
        names=["threshold", "fold", "seed"],
    )
    observed_threshold = pd.MultiIndex.from_frame(
        threshold[["threshold", "fold", "seed"]]
    )
    if len(observed_threshold) != 100 or set(observed_threshold) != set(
        expected_threshold
    ):
        add_mismatch(
            mismatches,
            location="threshold matrix",
            paper_value="5 thresholds x 4 folds x 5 seeds",
            computed_value=f"{len(observed_threshold)} unique keys",
            source="results/raw_runs.csv",
            reason="Threshold combinations are incomplete.",
        )

    required_run_files = (
        "config.json",
        "predictions.npz",
        "confusion_matrix.csv",
        "calibration_bins.csv",
        "training_history.csv",
    )
    missing_run_files = {}
    config_metadata_mismatches = []
    for row in raw.itertuples(index=False):
        run_dir = results_dir / "runs" / str(row.run_id)
        missing = [
            name for name in required_run_files if not (run_dir / name).is_file()
        ]
        if missing:
            missing_run_files[str(row.run_id)] = missing
            continue
        config = json.loads((run_dir / "config.json").read_text("utf-8"))
        recorded_run = config["run"]
        for key in ("fold", "seed", "method", "label_ratio", "threshold"):
            actual = recorded_run.get(key)
            expected_value = getattr(row, key)
            if actual is None and pd.isna(expected_value):
                continue
            if isinstance(actual, float) and isinstance(expected_value, float):
                if np.isclose(actual, expected_value, atol=1e-12):
                    continue
            if actual != expected_value:
                config_metadata_mismatches.append(
                    {"run_id": str(row.run_id), "field": key}
                )
                break
        if str(config["run"]["dataset_hash"]) != str(row.dataset_hash):
            config_metadata_mismatches.append(
                {"run_id": str(row.run_id), "field": "dataset_hash"}
            )
        if str(config["run"]["split_hash"]) != str(row.split_hash):
            config_metadata_mismatches.append(
                {"run_id": str(row.run_id), "field": "split_hash"}
            )
    if missing_run_files:
        add_mismatch(
            mismatches,
            location="per-run audit files",
            paper_value="config, predictions, confusion, calibration, and history",
            computed_value=f"{len(missing_run_files)} incomplete run directories",
            source="results/runs/",
            reason="One or more run-level audit files are missing.",
        )
    if config_metadata_mismatches:
        add_mismatch(
            mismatches,
            location="per-run configuration metadata",
            paper_value="matches raw_runs.csv",
            computed_value=f"{len(config_metadata_mismatches)} mismatches",
            source="results/runs/<run_id>/config.json",
            reason="Run configuration metadata differs from the canonical row.",
        )

    recheck_path = results_dir / "rechecks" / "all_runs.json"
    recheck = json.loads(recheck_path.read_text("utf-8"))
    recheck_passed = (
        int(recheck.get("passed", -1)) == len(raw)
        and len(recheck.get("runs", [])) == len(raw)
        and all(run.get("status") == "passed" for run in recheck["runs"])
        and all(
            all(item.get("passed") is True for item in run["metrics"].values())
            for run in recheck["runs"]
        )
    )
    if not recheck_passed:
        add_mismatch(
            mismatches,
            location="independent metric rechecks",
            paper_value="320 passed",
            computed_value=str(recheck.get("passed")),
            source="results/rechecks/all_runs.json",
            reason="Saved logits do not reproduce all recorded metrics.",
        )

    threshold_round_violations = int(
        threshold["self_training_rounds_completed"].fillna(0).gt(1).sum()
    )
    if threshold_round_violations:
        add_mismatch(
            mismatches,
            location="threshold self-training rounds",
            paper_value="at most one round",
            computed_value=f"{threshold_round_violations} violations",
            source="results/raw_runs.csv",
            reason="Threshold sensitivity used more than the controlled one round.",
        )

    temperature_metric_fields = [
        "macro_f1",
        "accuracy",
        "ece",
        "nll",
        "brier",
        "raw_macro_f1",
        "raw_accuracy",
        "raw_ece",
        "raw_nll",
        "raw_brier",
    ]
    temperature_identity = True
    for field in temperature_metric_fields:
        pivot = main.pivot_table(
            index=["fold", "seed"],
            columns="method",
            values=field,
        )
        if not np.allclose(
            pivot["cwt_gru"],
            pivot["cwt_gru_temp"],
            rtol=0,
            atol=0,
            equal_nan=True,
        ):
            temperature_identity = False
    if not temperature_identity:
        add_mismatch(
            mismatches,
            location="temperature-scaling evaluation",
            paper_value="same CWT-GRU weights and predictions",
            computed_value="metric difference found",
            source="results/raw_runs.csv",
            reason="Temperature scaling appears to have retrained or changed the model.",
        )

    dirty_rows = raw["git_dirty"].astype(str).str.casefold().isin(
        {"true", "1", "yes"}
    )
    return {
        "total_rows": int(len(raw)),
        "experiment_counts": {str(k): int(v) for k, v in counts.items()},
        "unique_run_ids": int(raw["run_id"].nunique()),
        "duplicate_run_ids": int(raw["run_id"].duplicated().sum()),
        "duplicate_experiment_keys": int(raw.duplicated(key_columns).sum()),
        "failed_runs": int((~raw["status"].eq("success")).sum()),
        "git_dirty_rows": int(dirty_rows.sum()),
        "missing_required_fields": null_counts,
        "main_matrix_complete": len(observed_main) == 140
        and set(observed_main) == set(expected_main),
        "label_matrix_complete": len(observed_label) == 80
        and set(observed_label) == set(expected_label),
        "threshold_matrix_complete": len(observed_threshold) == 100
        and set(observed_threshold) == set(expected_threshold),
        "missing_run_files": len(missing_run_files),
        "config_metadata_mismatches": len(config_metadata_mismatches),
        "independent_metric_rechecks_passed": recheck_passed,
        "threshold_round_violations": threshold_round_violations,
        "temperature_identity": temperature_identity,
    }


def check_external_completeness(
    external: pd.DataFrame,
    external_dir: Path,
    mismatches: list[dict[str, Any]],
) -> dict[str, Any]:
    methods = ("drsn", "tfnet")
    expected = pd.MultiIndex.from_product(
        [methods, FOLDS, SEEDS],
        names=["method", "fold", "seed"],
    )
    observed = pd.MultiIndex.from_frame(
        external[["method", "fold", "seed"]]
    )
    duplicate_keys = int(
        external.duplicated(["method", "fold", "seed"]).sum()
    )
    failed = int((~external["status"].eq("success")).sum())
    required = (
        "macro_f1",
        "accuracy",
        "best_val_macro_f1",
        "parameters",
        "gflops_sample",
        "train_time",
        "inference_ms_per_batch",
    )
    missing = {
        column: int(external[column].isna().sum())
        for column in required
        if external[column].isna().any()
    }
    if (
        len(external) != 40
        or len(observed) != 40
        or set(observed) != set(expected)
        or duplicate_keys
    ):
        add_mismatch(
            mismatches,
            location="external SOTA experiment matrix",
            paper_value="2 methods x 4 folds x 5 seeds",
            computed_value=f"{len(external)} rows, {len(observed)} unique keys",
            source="results/sota/raw_runs.csv",
            reason="External baseline matrix is incomplete or duplicated.",
        )
    if failed or missing:
        add_mismatch(
            mismatches,
            location="external SOTA result integrity",
            paper_value="40 successful complete rows",
            computed_value={"failed": failed, "missing": missing},
            source="results/sota/raw_runs.csv",
            reason="External baseline results contain failed or missing values.",
        )
    dirty = int(
        external["git_dirty"].astype(str).str.casefold().isin(
            {"true", "1", "yes"}
        ).sum()
    )
    if dirty:
        add_mismatch(
            mismatches,
            location="external SOTA provenance",
            paper_value="clean worktree",
            computed_value=f"{dirty} dirty runs",
            source="results/sota/raw_runs.csv",
            reason="External runs came from a dirty worktree.",
        )
    missing_files = []
    config_mismatches = []
    hostname_fields = 0
    required_files = (
        "config.json",
        "predictions.npz",
        "confusion_matrix.csv",
        "calibration_bins.csv",
        "training_history.csv",
    )
    for row in external.itertuples(index=False):
        run_dir = external_dir / "runs" / str(row.run_id)
        missing = [
            name for name in required_files if not (run_dir / name).is_file()
        ]
        if missing:
            missing_files.append({"run_id": str(row.run_id), "files": missing})
            continue
        config = json.loads((run_dir / "config.json").read_text("utf-8"))
        hostname_fields += int("hostname" in config.get("environment", {}))
        recorded_run = config["run"]
        for key in ("fold", "seed", "method", "label_ratio", "threshold"):
            actual = recorded_run.get(key)
            expected_value = getattr(row, key)
            if actual is None and pd.isna(expected_value):
                continue
            if isinstance(actual, float) and isinstance(expected_value, float):
                if np.isclose(actual, expected_value, atol=1e-12):
                    continue
            if actual != expected_value:
                config_mismatches.append(
                    {"run_id": str(row.run_id), "field": key}
                )
                break
    if missing_files or config_mismatches:
        add_mismatch(
            mismatches,
            location="external run artifacts",
            paper_value="complete config and diagnostics matching raw rows",
            computed_value={
                "missing": len(missing_files),
                "config_mismatches": len(config_mismatches),
            },
            source="results/sota/runs/",
            reason="External run-level audit material is incomplete.",
        )
    if hostname_fields:
        add_mismatch(
            mismatches,
            location="external run privacy metadata",
            paper_value="no hostname field",
            computed_value=f"{hostname_fields} files",
            source="results/sota/runs/<run_id>/config.json",
            reason="New external runs reintroduced host metadata.",
        )
    recheck_path = external_dir / "rechecks" / "all_runs.json"
    recheck = json.loads(recheck_path.read_text("utf-8"))
    recheck_ok = (
        int(recheck.get("passed", -1)) == len(external)
        and len(recheck.get("runs", [])) == len(external)
        and all(run.get("status") == "passed" for run in recheck["runs"])
        and all(
            all(item.get("passed") is True for item in run["metrics"].values())
            for run in recheck["runs"]
        )
    )
    if not recheck_ok:
        add_mismatch(
            mismatches,
            location="external SOTA metric rechecks",
            paper_value="40 passed",
            computed_value=str(recheck.get("passed")),
            source="results/sota/rechecks/all_runs.json",
            reason="External saved logits do not reproduce recorded metrics.",
        )
    return {
        "rows": int(len(external)),
        "group_counts": {
            str(key): int(value)
            for key, value in external.groupby("method").size().items()
        },
        "failed_runs": failed,
        "duplicate_keys": duplicate_keys,
        "dirty_runs": dirty,
        "missing_run_files": len(missing_files),
        "config_metadata_mismatches": len(config_mismatches),
        "hostname_fields": hostname_fields,
        "missing_required_fields": missing,
        "matrix_complete": len(observed) == 40 and set(observed) == set(expected),
        "independent_metric_rechecks_passed": recheck_ok,
    }


def check_paper_tables(
    raw: pd.DataFrame,
    document: Document,
    root: Path,
    results_dir: Path,
    mismatches: list[dict[str, Any]],
) -> dict[str, Any]:
    if len(document.tables) != 13:
        add_mismatch(
            mismatches,
            location="manuscript tables",
            paper_value="13",
            computed_value=str(len(document.tables)),
            source="docs/paper/Leakage_Safe_Bearing_Benchmark_SCI_final.docx",
            reason="Manuscript table count differs from the audited design.",
        )
    if len(document.inline_shapes) != 12:
        add_mismatch(
            mismatches,
            location="manuscript figures",
            paper_value="12",
            computed_value=str(len(document.inline_shapes)),
            source="docs/paper/Leakage_Safe_Bearing_Benchmark_SCI_final.docx",
            reason="Manuscript figure count differs from the audited design.",
        )

    main = raw.loc[raw["experiment_type"].eq("main")]
    config = json.loads((root / "configs" / "paderborn_main.json").read_text("utf-8"))
    metadata_path = (
        root / "data" / "public" / "paderborn" / "pu_4class_w8192.json"
    )
    metadata = (
        json.loads(metadata_path.read_text("utf-8"))
        if metadata_path.is_file()
        else None
    )
    recorded_run_config = json.loads(
        (root / str(main.iloc[0]["config_path"])).read_text("utf-8")
    )

    if len(document.tables) >= 1:
        labels = ["C0", "C1", "C2", "C3"]
        expected = [
            [
                "Fold",
                "Training conditions",
                "Validation",
                "Test",
                "Train windows",
                "Validation windows",
                "Test windows",
            ]
        ]
        for fold in FOLDS:
            manifest = json.loads(
                (results_dir / "splits" / f"fold{fold}.json").read_text(
                    "utf-8"
                )
            )
            expected.append(
                [
                    str(fold),
                    ", ".join(labels[index] for index in manifest["train_conditions"]),
                    labels[int(manifest["validation_condition"])],
                    labels[int(manifest["test_condition"])],
                    str(manifest["train"]["count"]),
                    str(manifest["validation"]["count"]),
                    str(manifest["test"]["count"]),
                ]
            )
        compare_table(
            table_rows(document.tables[0]),
            expected,
            location="Table 1",
            source="results/splits/fold*.json",
            mismatches=mismatches,
        )

    if len(document.tables) >= 2:
        expected = [
            ["Module", "Operation in code", "Output shape"],
            ["CWT input", "Three channel magnitude maps", "3 x 96 x 96"],
            [
                "Stem",
                "Conv5x5(32, stride 2) + BN + SiLU + MaxPool2",
                "32 x 24 x 24",
            ],
            ["Residual stage 1", "Two Conv3x3 + BN + optional SE", "64 x 24 x 24"],
            [
                "Residual stage 2",
                "Two Conv3x3, stride 2 + BN + optional SE",
                "128 x 12 x 12",
            ],
            [
                "Residual stage 3",
                "Two Conv3x3, stride 2 + BN + optional SE",
                "192 x 6 x 6",
            ],
            ["Image feature", "Global average pooling", "192"],
            ["Raw sequence", "Transpose 3 x 8192 to 8192 x 3", "8192 x 3"],
            ["Recurrent feature", "Two-layer bidirectional GRU, hidden 64", "128"],
            ["Fusion classifier", "320 -> 256 -> 128 -> 4", "4 logits"],
        ]
        compare_table(
            table_rows(document.tables[1]),
            expected,
            location="Table 2",
            source="src/spindle_cnn/models/*.py and data metadata",
            mismatches=mismatches,
        )
        if metadata is not None:
            raw_shape = metadata["shape"]["raw"]
            cwt_shape = metadata["shape"]["cwt"]
            if raw_shape[1:] != [3, 8192] or cwt_shape[1:] != [3, 96, 96]:
                add_mismatch(
                    mismatches,
                    location="Table 2 input shapes",
                    paper_value="3 x 8192 raw and 3 x 96 x 96 CWT",
                    computed_value=f"{raw_shape[1:]} and {cwt_shape[1:]}",
                    source=str(metadata_path.relative_to(root)),
                    reason="Cached tensor shapes differ from the manuscript.",
                )

    if len(document.tables) >= 3:
        training = config["training"]
        expected = [
            ["Parameter", "Value"],
            ["Optimizer", "AdamW [19,20]"],
            ["Initial learning rate", f"{float(training['learning_rate']):g}"],
            ["Weight decay", f"{float(training['weight_decay']):g}"],
            ["Batch size", str(training["batch_size"])],
            ["Epochs", f"maximum {training['epochs']}"],
            [
                "Early stopping",
                f"validation macro-F1, patience {training['patience']}",
            ],
            ["Scheduler", f"cosine annealing over {training['epochs']} epochs"],
            ["Precision", "mixed precision on CUDA"],
            ["Gradient clipping", "maximum norm 5.0"],
            [
                "Input augmentation",
                "raw circular shift up to +/-32 samples and corresponding CWT shift; "
                "50% probability of additive noise",
            ],
            [
                "GPU",
                recorded_run_config["environment"]["device_name"],
            ],
        ]
        compare_table(
            table_rows(document.tables[2]),
            expected,
            location="Table 3",
            source="configs/paderborn_main.json and per-run config",
            mismatches=mismatches,
        )

    if len(document.tables) >= 4:
        counts = raw.groupby("experiment_type").size()
        threshold_count = int(
            raw.loc[
                raw["experiment_type"].eq("threshold"), "threshold"
            ].nunique()
        )
        expected = [
            ["Experiment", "Purpose", "Runs"],
            [
                "Main",
                f"Seven method conditions; {len(FOLDS)} folds x {len(SEEDS)} seeds; "
                "Full balanced budget",
                str(int(counts["main"])),
            ],
            [
                "Label efficiency",
                "Supervised and self-training at 10% balanced budget and "
                "50% balanced budget; Full balanced budget reuses main rows",
                str(int(counts["label_efficiency"])),
            ],
            [
                "Threshold sensitivity",
                f"{threshold_count} thresholds at 50% balanced budget; "
                f"{len(FOLDS)} folds x {len(SEEDS)} seeds",
                str(int(counts["threshold"])),
            ],
            ["Total", "Canonical successful evaluations", str(len(raw))],
        ]
        compare_table(
            table_rows(document.tables[3]),
            expected,
            location="Table 4",
            source="results/raw_runs.csv and configs/paderborn_main.json",
            mismatches=mismatches,
        )

    summary = independent_summary_by(raw, "main", "method")
    if len(document.tables) >= 5:
        expected = [["Method", "Macro-F1", "Accuracy", "Runs"]]
        for method in METHOD_ORDER:
            row = summary.loc[method]
            expected.append(
                [
                    METHOD_LABELS[method],
                    mean_sd(row["mean"], row["std"]),
                    mean_sd(row["accuracy_mean"], row["accuracy_std"]),
                    str(int(row["runs"])),
                ]
            )
        compare_table(
            table_rows(document.tables[4]),
            expected,
            location="Table 5",
            source="results/raw_runs.csv",
            mismatches=mismatches,
        )

    if len(document.tables) >= 6:
        hierarchical = pd.read_csv(
            results_dir / "hierarchical_statistics.csv"
        )
        expected = [
            [
                "Method",
                "Paired delta Macro-F1",
                "Hierarchical 95% interval",
                "Fold-aware p",
                "Stable direction?",
            ]
        ]
        for method in METHOD_ORDER[1:]:
            comparison = f"main:raw_cnn:{method}"
            row = hierarchical.loc[
                hierarchical["comparison"].eq(comparison)
            ].iloc[0]
            expected.append(
                [
                    METHOD_LABELS[method],
                    f"{float(row['mean_delta']):+.3f}",
                    f"{float(row['hierarchical_ci_low']):+.3f} to "
                    f"{float(row['hierarchical_ci_high']):+.3f}",
                    f"p {p_text(row['fold_aware_p'])}",
                    (
                        "Yes"
                        if bool(row["hierarchical_ci_excludes_zero"])
                        else "No"
                    ),
                ]
            )
        compare_table(
            table_rows(document.tables[5]),
            expected,
            location="Table 6",
            source="results/hierarchical_statistics.csv",
            mismatches=mismatches,
        )

    label_frame = frame_with_label_100(raw)
    label = label_frame.loc[label_frame["experiment_type"].eq("label_efficiency")]
    if len(document.tables) >= 7:
        expected = [["Label budget", "Method", "Macro-F1", "Accuracy", "Runs"]]
        for ratio in (0.1, 0.5, 1.0):
            for method, label_text in (
                ("cwt_gru", "Supervised"),
                ("cwt_gru_selftrain", "Calibrated self-training"),
            ):
                group = label.loc[
                    label["method"].eq(method) & label["label_ratio"].eq(ratio)
                ]
                expected.append(
                    [
                        {
                            0.1: "10% balanced budget",
                            0.5: "50% balanced budget",
                            1.0: "Full balanced budget",
                        }[ratio],
                        label_text,
                        mean_sd(group["macro_f1"].mean(), group["macro_f1"].std()),
                        mean_sd(group["accuracy"].mean(), group["accuracy"].std()),
                        str(len(group)),
                    ]
                )
        compare_table(
            table_rows(document.tables[6]),
            expected,
            location="Table 7",
            source="results/raw_runs.csv",
            mismatches=mismatches,
        )

    threshold = raw.loc[raw["experiment_type"].eq("threshold")]
    if len(document.tables) >= 8:
        expected = [
            [
                "Threshold",
                "Mean pseudo-labels",
                "Coverage",
                "Pseudo-label accuracy",
                "Test Macro-F1",
                "Runs",
            ]
        ]
        for value in (0.3, 0.4, 0.5, 0.6, 0.7):
            group = threshold.loc[threshold["threshold"].eq(value)]
            expected.append(
                [
                    f"{value:.1f}",
                    f"{group['pseudo_count'].mean():.1f}",
                    percent(group["pseudo_coverage"].mean()),
                    f"{group['pseudo_accuracy'].mean():.3f}",
                    mean_sd(group["macro_f1"].mean(), group["macro_f1"].std()),
                    str(len(group)),
                ]
            )
        compare_table(
            table_rows(document.tables[7]),
            expected,
            location="Table 8",
            source="results/raw_runs.csv",
            mismatches=mismatches,
        )

    if len(document.tables) >= 9:
        expected = [
            [
                "Method",
                "ECE raw",
                "ECE scaled",
                "NLL raw",
                "NLL scaled",
                "Brier raw",
                "Brier scaled",
            ]
        ]
        for method in ("raw_cnn", "cwt_cnn", "cwt_gru", "cwt_gru_selftrain"):
            group = main.loc[main["method"].eq(method)]
            expected.append(
                [
                    METHOD_LABELS[method],
                    f"{group['raw_ece'].mean():.3f}",
                    f"{group['ece'].mean():.3f}",
                    f"{group['raw_nll'].mean():.3f}",
                    f"{group['nll'].mean():.3f}",
                    f"{group['raw_brier'].mean():.3f}",
                    f"{group['brier'].mean():.3f}",
                ]
            )
        compare_table(
            table_rows(document.tables[8]),
            expected,
            location="Table 9",
            source="results/raw_runs.csv",
            mismatches=mismatches,
        )

    if len(document.tables) >= 10:
        expected = [
            [
                "Method",
                "Input representation",
                "Core architecture",
                "Parameters",
                "Macro-F1",
                "Accuracy",
                "Runs",
            ]
        ]
        architecture = {
            "raw_cnn": ("Raw waveform", "Three-layer 1D CNN + linear head"),
            "gru_only": (
                "Raw waveform",
                "Two-layer bidirectional GRU + linear head",
            ),
            "cwt_cnn": (
                "CWT magnitude maps",
                "Residual CNN with SE + linear head",
            ),
            "cwt_gru_no_se": (
                "Raw + CWT",
                "Residual image CNN + two-layer bidirectional GRU",
            ),
            "cwt_gru": (
                "Raw + CWT",
                "Residual CNN with SE + two-layer bidirectional GRU",
            ),
            "cwt_gru_temp": (
                "Raw + CWT",
                "CWT-GRU weights; validation temperature scaling",
            ),
            "cwt_gru_selftrain": (
                "Raw + CWT",
                "CWT-GRU; thresholded pseudo-label retraining",
            ),
        }
        for method in METHOD_ORDER:
            row = summary.loc[method]
            representation, core = architecture[method]
            expected.append(
                [
                    METHOD_LABELS[method],
                    representation,
                    core,
                    f"{int(row['parameters']):,}",
                    mean_sd(row["mean"], row["std"]),
                    mean_sd(row["accuracy_mean"], row["accuracy_std"]),
                    str(int(row["runs"])),
                ]
            )
        external_raw = pd.read_csv(results_dir / "sota" / "raw_runs.csv")
        external_summary = pd.read_csv(
            results_dir / "sota" / "summary_with_ci.csv"
        )
        external_architecture = {
            "drsn": (
                "Raw waveform",
                "Four-stage residual shrinkage CNN; adaptive soft thresholding",
            ),
            "tfnet": (
                "Raw waveform",
                "Learnable sinusoidal TF convolution + CNN backbone",
            ),
        }
        external_labels = {
            "drsn": "DRSN (same-protocol adaptation)",
            "tfnet": "TFN (same-protocol adaptation)",
        }
        for method, (representation, core) in external_architecture.items():
            row = external_summary.loc[external_summary["Method"].eq(method)].iloc[0]
            parameters = int(
                external_raw.loc[
                    external_raw["method"].eq(method), "parameters"
                ].median()
            )
            expected.append(
                [
                    external_labels[method],
                    representation,
                    core,
                    f"{parameters:,}",
                    mean_sd(row["Macro-F1 Mean"], row["Macro-F1 Std"]),
                    mean_sd(row["Accuracy Mean"], row["Accuracy Std"]),
                    str(int(row["Runs"])),
                ]
            )
        compare_table(
            table_rows(document.tables[9]),
            expected,
            location="Table 10",
            source="results/raw_runs.csv",
            mismatches=mismatches,
        )

    components = {
        "raw_cnn": (0, 0, 0, 0, 0),
        "gru_only": (0, 1, 0, 0, 0),
        "cwt_cnn": (1, 0, 1, 0, 0),
        "cwt_gru_no_se": (1, 1, 0, 0, 0),
        "cwt_gru": (1, 1, 1, 0, 0),
        "cwt_gru_temp": (1, 1, 1, 1, 0),
        "cwt_gru_selftrain": (1, 1, 1, 1, 1),
    }
    if len(document.tables) >= 11:
        expected = [
            [
                "Model",
                "CWT",
                "GRU",
                "SE",
                "Calibration",
                "Self-training",
                "Macro-F1",
                "Accuracy",
                "Runs",
            ]
        ]
        for method in METHOD_ORDER:
            row = summary.loc[method]
            expected.append(
                [
                    METHOD_LABELS[method],
                    *[str(value) for value in components[method]],
                    mean_sd(row["mean"], row["std"]),
                    mean_sd(row["accuracy_mean"], row["accuracy_std"]),
                    str(int(row["runs"])),
                ]
            )
        compare_table(
            table_rows(document.tables[10]),
            expected,
            location="Table 11",
            source="results/raw_runs.csv",
            mismatches=mismatches,
        )

    if len(document.tables) >= 12:
        expected = [
            [
                "Method",
                "Parameters",
                "GFLOPs/sample",
                "Training time (s)",
                "Inference (ms/batch)",
            ]
        ]
        for method in (
            "raw_cnn",
            "gru_only",
            "cwt_cnn",
            "cwt_gru_no_se",
            "cwt_gru",
            "cwt_gru_selftrain",
        ):
            group = main.loc[main["method"].eq(method)]
            expected.append(
                [
                    METHOD_LABELS[method],
                    f"{int(group['parameters'].median()):,}",
                    f"{group['gflops_sample'].median():.3f}",
                    f"{group['train_time'].median():.1f}",
                    f"{group['inference_ms_per_batch'].median():.1f}",
                ]
            )
        compare_table(
            table_rows(document.tables[11]),
            expected,
            location="Table 12",
            source="results/raw_runs.csv",
            mismatches=mismatches,
        )

    if len(document.tables) >= 13:
        domain_gap = pd.read_csv(results_dir / "domain_gap_summary.csv")
        expected = [
            [
                "Model",
                "Validation Macro-F1",
                "Test Macro-F1",
                "Absolute gap",
                "Relative drop",
                "Runs",
            ]
        ]
        for row in domain_gap.itertuples(index=False):
            method = str(row.Model)
            expected.append(
                [
                    METHOD_LABELS[method],
                    mean_sd(row._1, row._2),
                    mean_sd(row._3, row._4),
                    f"{float(row._5):.3f}",
                    f"{float(row._7):.1f}%",
                    str(int(row.Runs)),
                ]
            )
        compare_table(
            table_rows(document.tables[12]),
            expected,
            location="Table 13",
            source="results/domain_gap_summary.csv",
            mismatches=mismatches,
        )

    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    main_pivot = main.pivot_table(
        index=["fold", "seed"],
        columns="method",
        values="macro_f1",
    )
    no_se_delta = (
        main_pivot["cwt_gru_no_se"] - main_pivot["cwt_gru"]
    )
    label_50 = label_frame.loc[
        label_frame["experiment_type"].eq("label_efficiency")
        & label_frame["label_ratio"].eq(0.5)
    ]
    label_50_pivot = label_50.pivot_table(
        index=["fold", "seed"],
        columns="method",
        values="macro_f1",
    )
    label_50_delta = (
        label_50_pivot["cwt_gru_selftrain"] - label_50_pivot["cwt_gru"]
    )
    label_100 = label_frame.loc[
        label_frame["experiment_type"].eq("label_efficiency")
        & label_frame["label_ratio"].eq(1.0)
    ]
    label_100_pivot = label_100.pivot_table(
        index=["fold", "seed"],
        columns="method",
        values="macro_f1",
    )
    label_100_delta = (
        label_100_pivot["cwt_gru_selftrain"] - label_100_pivot["cwt_gru"]
    )
    threshold_summary = (
        threshold.groupby("threshold")
        .agg(macro_f1=("macro_f1", "mean"), coverage=("pseudo_coverage", "mean"))
        .sort_index()
    )
    confusion_frames = [
        pd.read_csv(results_dir / "runs" / run_id / "confusion_matrix.csv")
        for run_id in main.loc[main["method"].eq("cwt_gru"), "run_id"]
    ]
    confusion_counts = pd.concat(confusion_frames, ignore_index=True).groupby(
        ["true_class", "predicted_class"], as_index=False
    )["count"].sum()
    confusion = confusion_counts.pivot(
        index="true_class",
        columns="predicted_class",
        values="count",
    ).reindex(index=range(4), columns=range(4), fill_value=0)
    confusion = confusion.div(confusion.sum(axis=1), axis=0)
    cwt_calibration = main.loc[main["method"].eq("cwt_gru")]
    external_summary = pd.read_csv(
        results_dir / "sota" / "summary_with_ci.csv"
    )
    external_tests = pd.read_csv(
        results_dir / "sota" / "statistical_tests.csv"
    )
    complexity = main.groupby("method").agg(
        parameters=("parameters", "median"),
        train_time=("train_time", "median"),
    )
    top_threshold = threshold_summary["macro_f1"].idxmax()
    required_fragments = {
        "Title": "Leakage-Safe Benchmarking of Time-Frequency Fusion",
        "RQ1": "under leakage-safe cross-domain evaluation?",
        "RQ6": "only confidence reliability?",
        "Raw result": mean_sd(
            summary.loc["raw_cnn", "mean"], summary.loc["raw_cnn", "std"]
        ),
        "CWT result": mean_sd(
            summary.loc["cwt_cnn", "mean"], summary.loc["cwt_cnn", "std"]
        ),
        "CWT-GRU result": mean_sd(
            summary.loc["cwt_gru", "mean"], summary.loc["cwt_gru", "std"]
        ),
        "No-SE sensitivity p-value": (
            f"p {p_text(stats.wilcoxon(no_se_delta).pvalue)}"
        ),
        "50% self-training result": (
            f"paired difference was {label_50_delta.mean():+.3f}, with a "
            "fold-aware interval from"
        ),
        "Full balanced self-training result": (
            f"paired difference was {label_100_delta.mean():+.3f}, with interval"
        ),
        "Threshold minimum": f"{threshold_summary['macro_f1'].min():.3f}",
        "Threshold maximum": f"{threshold_summary['macro_f1'].max():.3f}",
        "Threshold accuracy at 0.3": (
            f"{threshold.loc[threshold['threshold'].eq(0.3), 'pseudo_accuracy'].mean():.3f}"
        ),
        "Threshold accuracy at 0.7": (
            f"{threshold.loc[threshold['threshold'].eq(0.7), 'pseudo_accuracy'].mean():.3f}"
        ),
        "Threshold coverage at 0.3": percent(
            threshold.loc[threshold["threshold"].eq(0.3), "pseudo_coverage"].mean()
        ),
        "Threshold coverage at 0.7": percent(
            threshold.loc[threshold["threshold"].eq(0.7), "pseudo_coverage"].mean()
        ),
        "Highest threshold result": f"tau = {top_threshold:.1f}",
        "Label 50% means": (
            f"{label_50.loc[label_50['method'].eq('cwt_gru'), 'macro_f1'].mean():.3f} "
            f"to {label_50.loc[label_50['method'].eq('cwt_gru_selftrain'), 'macro_f1'].mean():.3f}"
        ),
        "CWT-GRU ECE": (
            f"{cwt_calibration['raw_ece'].mean():.3f} to "
            f"{cwt_calibration['ece'].mean():.3f}"
        ),
        "CWT-GRU NLL": (
            f"{cwt_calibration['raw_nll'].mean():.3f} to "
            f"{cwt_calibration['nll'].mean():.3f}"
        ),
        "CWT-GRU inner-ring recall": percent(confusion.loc[2, 2]),
        "CWT-GRU outer-ring recall": percent(confusion.loc[1, 1]),
        "CWT-GRU outer-to-inner confusion": percent(confusion.loc[1, 2]),
        "Parameter ratio": (
            f"{complexity.loc['cwt_gru', 'parameters'] / complexity.loc['raw_cnn', 'parameters']:.1f} times"
        ),
        "CWT-GRU median training time": (
            f"{complexity.loc['cwt_gru', 'train_time']:.1f} s"
        ),
        "Raw CNN median training time": (
            f"{complexity.loc['raw_cnn', 'train_time']:.1f} s"
        ),
        "DRSN result": mean_sd(
            external_summary.loc[
                external_summary["Method"].eq("drsn"), "Macro-F1 Mean"
            ].iloc[0],
            external_summary.loc[
                external_summary["Method"].eq("drsn"), "Macro-F1 Std"
            ].iloc[0],
        ),
        "TFN result": mean_sd(
            external_summary.loc[
                external_summary["Method"].eq("tfnet"), "Macro-F1 Mean"
            ].iloc[0],
            external_summary.loc[
                external_summary["Method"].eq("tfnet"), "Macro-F1 Std"
            ].iloc[0],
        ),
        "DRSN paired delta": (
            f"{float(hierarchical.loc[hierarchical['comparison'].eq('external:raw_cnn:drsn'), 'mean_delta'].iloc[0]):+.3f}"
        ),
        "TFN paired delta": (
            f"{float(hierarchical.loc[hierarchical['comparison'].eq('external:raw_cnn:tfnet'), 'mean_delta'].iloc[0]):+.3f}"
        ),
        "Domain gap statement": (
            "cannot be explained merely by an inability to fit the development domains"
        ),
        "Contribution wording": (
            "joint comparison that distinguishes representation and training"
        ),
        "SOTA boundary": "No unreproduced external result is included.",
        "TFN reference DOI": "10.1016/j.ymssp.2023.110952",
        "Spindle scope": "not a complete machine-tool spindle assembly",
    }
    for label, fragment in required_fragments.items():
        if fragment not in text:
            add_mismatch(
                mismatches,
                location=f"manuscript text: {label}",
                paper_value=fragment,
                computed_value="fragment not found",
                source="docs/paper/Leakage_Safe_Bearing_Benchmark_SCI_final.docx",
                reason="A required result statement is absent or inconsistent.",
            )
    for table_number in range(1, 14):
        if f"Table {table_number}." not in text:
            add_mismatch(
                mismatches,
                location=f"Table {table_number} caption",
                paper_value=f"Table {table_number}.",
                computed_value="caption not found",
                source="docs/paper/Leakage_Safe_Bearing_Benchmark_SCI_final.docx",
                reason="A numbered table caption is missing.",
            )
    for figure_number in range(1, 13):
        if f"Figure {figure_number}." not in text:
            add_mismatch(
                mismatches,
                location=f"Figure {figure_number} caption",
                paper_value=f"Figure {figure_number}.",
                computed_value="caption not found",
                source="docs/paper/Leakage_Safe_Bearing_Benchmark_SCI_final.docx",
                reason="A numbered figure caption is missing.",
            )
    if "spindle bearing" in text.casefold():
        add_mismatch(
            mismatches,
            location="manuscript scope terminology",
            paper_value="no unsupported spindle-bearing scope claim",
            computed_value="unsupported phrase found",
            source="docs/paper/Leakage_Safe_Bearing_Benchmark_SCI_final.docx",
            reason="The manuscript overstates the evidence from a bearing test rig.",
        )
    for obsolete in (
        "TFNet",
        "Full-label",
        "full-label",
        "100% label",
        "p=<",
        "p =<",
        "training failure",
    ):
        if obsolete.casefold() in text.casefold():
            add_mismatch(
                mismatches,
                location="balanced-budget terminology",
                paper_value=(
                    "10% balanced budget, 50% balanced budget, "
                    "Full balanced budget; TFN; standard p-value syntax"
                ),
                computed_value=obsolete,
                source="docs/paper/Leakage_Safe_Bearing_Benchmark_SCI_final.docx",
                reason="Obsolete terminology or p-value formatting remains.",
            )

    return {
        "tables": len(document.tables),
        "figures": len(document.inline_shapes),
        "required_text_fragments_checked": len(required_fragments),
    }


def independent_summary_by(
    raw: pd.DataFrame,
    experiment_type: str,
    group_column: str,
) -> pd.DataFrame:
    frame = raw.loc[raw["experiment_type"].eq(experiment_type)]
    rows = []
    for key, group in frame.groupby(group_column, sort=False):
        rows.append(
            {
                group_column: key,
                "runs": len(group),
                "mean": group["macro_f1"].mean(),
                "std": group["macro_f1"].std(ddof=1),
                "accuracy_mean": group["accuracy"].mean(),
                "accuracy_std": group["accuracy"].std(ddof=1),
                "parameters": group["parameters"].median(),
            }
        )
    return pd.DataFrame(rows).set_index(group_column)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--root",
        type=Path,
        default=Path(__file__).resolve().parents[1],
    )
    parser.add_argument(
        "--results-dir",
        type=Path,
        help="Results directory; defaults to <root>/results.",
    )
    parser.add_argument(
        "--external-dir",
        type=Path,
        help="External results directory; defaults to <results-dir>/sota.",
    )
    parser.add_argument(
        "--paper",
        type=Path,
        help=(
            "Generated manuscript DOCX; defaults to "
            "<root>/docs/paper/Leakage_Safe_Bearing_Benchmark_SCI_final.docx."
        ),
    )
    parser.add_argument("--json-output", type=Path)
    args = parser.parse_args()
    root = args.root.resolve()
    results_dir = (
        args.results_dir
        if args.results_dir and args.results_dir.is_absolute()
        else root / args.results_dir
        if args.results_dir
        else root / "results"
    ).resolve()
    external_dir = (
        args.external_dir
        if args.external_dir and args.external_dir.is_absolute()
        else root / args.external_dir
        if args.external_dir
        else results_dir / "sota"
    ).resolve()
    paper_path = (
        args.paper
        if args.paper and args.paper.is_absolute()
        else root / args.paper
        if args.paper
        else (
            root
            / "docs"
            / "paper"
            / "Leakage_Safe_Bearing_Benchmark_SCI_final.docx"
        )
    ).resolve()
    raw = pd.read_csv(results_dir / "raw_runs.csv")
    external = pd.read_csv(external_dir / "raw_runs.csv")
    summary = pd.read_csv(results_dir / "summary.csv")
    summary_ci = pd.read_csv(results_dir / "summary_with_ci.csv")
    pairs = pd.read_csv(results_dir / "statistical_tests.csv")
    label_summary = pd.read_csv(
        results_dir / "label_efficiency_summary.csv"
    )
    threshold_summary = pd.read_csv(results_dir / "threshold_summary.csv")
    ablation_summary = pd.read_csv(results_dir / "ablation_summary.csv")
    model_complexity = pd.read_csv(results_dir / "model_complexity.csv")
    hierarchical_statistics = pd.read_csv(
        results_dir / "hierarchical_statistics.csv"
    )
    mismatches: list[dict[str, Any]] = []

    completeness = check_run_completeness(raw, results_dir, mismatches)
    external_completeness = check_external_completeness(
        external,
        external_dir,
        mismatches,
    )
    check_summary_sources(
        raw,
        summary,
        summary_ci,
        pairs,
        label_summary,
        threshold_summary,
        ablation_summary,
        model_complexity,
        mismatches,
    )
    check_hierarchical_statistics(
        raw,
        external,
        hierarchical_statistics,
        mismatches,
    )

    if paper_path.is_file():
        paper = check_paper_tables(
            raw,
            Document(paper_path),
            root,
            results_dir,
            mismatches,
        )
    else:
        add_mismatch(
            mismatches,
            location="manuscript availability",
            paper_value="generated manuscript DOCX",
            computed_value=f"missing: {paper_path}",
            source=str(paper_path),
            reason="The manuscript must be built before numerical table audit.",
        )
        paper = {"available": False}

    payload = {
        "status": "PASS" if not mismatches else "FAIL",
        "raw_sha256": __import__("hashlib").sha256(
            (results_dir / "raw_runs.csv").read_bytes()
        ).hexdigest(),
        "results_dir": str(results_dir),
        "external_dir": str(external_dir),
        "paper_path": str(paper_path),
        "completeness": completeness,
        "external_completeness": external_completeness,
        "paper": paper,
        "mismatch_count": len(mismatches),
        "mismatches": mismatches,
    }
    if args.json_output:
        args.json_output.parent.mkdir(parents=True, exist_ok=True)
        args.json_output.write_text(
            json.dumps(payload, indent=2, ensure_ascii=False) + "\n",
            encoding="utf-8",
        )
    print(json.dumps(payload, indent=2, ensure_ascii=False))
    raise SystemExit(0 if not mismatches else 1)


if __name__ == "__main__":
    main()
